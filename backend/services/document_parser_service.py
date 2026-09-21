from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_DOCUMENT_BYTES = 10_000_000

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    file_type: str
    text: str


class DocumentParserService:
    """
    Extract normalized plain text from supported document formats.

    This service intentionally handles parsing only. Persistence,
    chunking, embeddings, ownership, and retrieval remain inside
    KnowledgeService and the knowledge API boundary.
    """

    def parse(
        self,
        *,
        filename: str,
        content: bytes,
    ) -> ParsedDocument:
        normalized_filename = Path(
            str(filename).strip()
        ).name

        if not normalized_filename:
            raise ValueError(
                "filename cannot be empty."
            )

        extension = Path(
            normalized_filename
        ).suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                "Unsupported document type. "
                "Only PDF and DOCX files are supported."
            )

        if not content:
            raise ValueError(
                "Document content cannot be empty."
            )

        if len(content) > MAX_DOCUMENT_BYTES:
            raise ValueError(
                "Document exceeds the maximum upload size."
            )

        try:
            if extension == ".pdf":
                text = self._parse_pdf(
                    content
                )
                file_type = "pdf"
            else:
                text = self._parse_docx(
                    content
                )
                file_type = "docx"
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(
                "The document could not be parsed."
            ) from exc

        normalized_text = self._normalize_text(
            text
        )

        if not normalized_text:
            raise ValueError(
                "The document does not contain usable text."
            )

        return ParsedDocument(
            filename=normalized_filename,
            file_type=file_type,
            text=normalized_text,
        )

    @staticmethod
    def _normalize_text(
        value: str,
    ) -> str:
        blocks: list[str] = []

        for raw_line in str(value).splitlines():
            line = " ".join(
                raw_line.split()
            ).strip()

            if line:
                blocks.append(line)

        return "\n\n".join(blocks)

    @staticmethod
    def _parse_pdf(
        content: bytes,
    ) -> str:
        reader = PdfReader(
            BytesIO(content)
        )

        page_text: list[str] = []

        for page in reader.pages:
            extracted = page.extract_text() or ""

            if extracted.strip():
                page_text.append(extracted)

        return "\n\n".join(page_text)

    @staticmethod
    def _parse_docx(
        content: bytes,
    ) -> str:
        document = Document(
            BytesIO(content)
        )

        parts: list[str] = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(
                    paragraph.text
                )

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                cells = [
                    cell
                    for cell in cells
                    if cell
                ]

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

        return "\n".join(parts)
