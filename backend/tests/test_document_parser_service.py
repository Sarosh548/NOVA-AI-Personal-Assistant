from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

import services.document_parser_service as parser_module
from services.document_parser_service import (
    MAX_DOCUMENT_BYTES,
    DocumentParserService,
)


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph(
        "NOVA is a personal AI assistant."
    )
    document.add_paragraph(
        "It uses memory and knowledge retrieval."
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "NOVA"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Active"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakePdfPage:
    def __init__(self, text: str):
        self.text = text

    def extract_text(self):
        return self.text


class FakePdfReader:
    def __init__(self, stream):
        self.pages = [
            FakePdfPage("NOVA\nAI assistant."),
            FakePdfPage("Knowledge retrieval\nworks."),
        ]


def test_parse_docx_extracts_paragraphs_and_tables():
    service = DocumentParserService()

    result = service.parse(
        filename="nova-guide.docx",
        content=make_docx(),
    )

    assert result.filename == "nova-guide.docx"
    assert result.file_type == "docx"
    assert (
        "NOVA is a personal AI assistant."
        in result.text
    )
    assert (
        "It uses memory and knowledge retrieval."
        in result.text
    )
    assert "Project | NOVA" in result.text
    assert "Status | Active" in result.text


def test_parse_pdf_extracts_page_text(
    monkeypatch,
):
    monkeypatch.setattr(
        parser_module,
        "PdfReader",
        FakePdfReader,
    )

    service = DocumentParserService()

    result = service.parse(
        filename="nova-guide.pdf",
        content=b"fake pdf bytes",
    )

    assert result.filename == "nova-guide.pdf"
    assert result.file_type == "pdf"
    assert result.text == (
        "NOVA\n\nAI assistant.\n\n"
        "Knowledge retrieval\n\nworks."
    )


@pytest.mark.parametrize(
    "filename",
    [
        "",
        "notes.txt",
        "document.doc",
        "image.png",
    ],
)
def test_parse_rejects_unsupported_or_empty_filename(
    filename,
):
    service = DocumentParserService()

    with pytest.raises(ValueError):
        service.parse(
            filename=filename,
            content=b"content",
        )


def test_parse_rejects_empty_content():
    service = DocumentParserService()

    with pytest.raises(
        ValueError,
        match="cannot be empty",
    ):
        service.parse(
            filename="empty.docx",
            content=b"",
        )


def test_parse_rejects_oversized_content():
    service = DocumentParserService()

    with pytest.raises(
        ValueError,
        match="maximum upload size",
    ):
        service.parse(
            filename="large.docx",
            content=b"x"
            * (MAX_DOCUMENT_BYTES + 1),
        )


def test_parse_rejects_password_or_corrupt_pdf(
    monkeypatch,
):
    class FailingPdfReader:
        def __init__(self, stream):
            raise RuntimeError(
                "invalid pdf"
            )

    monkeypatch.setattr(
        parser_module,
        "PdfReader",
        FailingPdfReader,
    )

    service = DocumentParserService()

    with pytest.raises(
        ValueError,
        match="could not be parsed",
    ):
        service.parse(
            filename="broken.pdf",
            content=b"not a real pdf",
        )


def test_normalize_text_collapses_whitespace():
    result = (
        DocumentParserService._normalize_text(
            "  Hello   NOVA  \n\n"
            "  Knowledge\tBase  "
        )
    )

    assert result == (
        "Hello NOVA\n\nKnowledge Base"
    )
